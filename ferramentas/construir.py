"""Constrói o site (docs/), o plano de leitura em PDF e a planilha de acompanhamento a partir de data/catalogo.json.

Uso:  python3 ferramentas/construir.py
Requisitos: Python 3.10+, weasyprint (PDF) e openpyxl (planilha). O site em si é HTML/CSS/JS puro.
"""
import csv, html, json, os, sys
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')
SITE = 'https://oaleffemanuel.github.io/biblioteca-italiana/'
REPO = 'https://github.com/oaleffemanuel/biblioteca-italiana'
PDF_NAME = 'plano-de-leitura-biblioteca-italiana.pdf'
XLSX_NAME = 'diario-de-leitura-biblioteca-italiana.xlsx'
TODAY = date.today().isoformat()
import hashlib
VER = hashlib.sha1(b''.join(open(os.path.join(DOCS, 'assets', f), 'rb').read() for f in ('style.css', 'app.js'))).hexdigest()[:8]

W = json.load(open(os.path.join(ROOT, 'data', 'catalogo.json'), encoding='utf-8'))
for i, w in enumerate(W): w['i'] = i
e = lambda s: html.escape(str(s or ''), quote=True)

# ───────────────────────── o percurso: 7 tappe ─────────────────────────
TAPPE = {
    1: dict(nome='Primi passi', pt='Primeiros passos', nivel='A1',
            curta='Histórias de 3 a 5 páginas, frases curtas e vocabulário do dia a dia. Quase todas com áudio.',
            perfil='Você conhece o presente do indicativo, se apresenta e entende frases simples.',
            objetivo='Criar o hábito de ler e se acostumar com o som do italiano escrito.',
            como='Leia cada história três vezes: só lendo; lendo e ouvindo; só ouvindo. Anote no máximo cinco palavras por história.',
            meta='As 7 histórias em 1 a 2 semanas.',
            avancar='Você entende a história inteira na segunda leitura, sem dicionário.'),
    2: dict(nome='Prendere il ritmo', pt='Pegando o ritmo', nivel='A1–A2',
            curta='Contos curtos de suspense, humor e cotidiano. 18 deles com áudio gravado.',
            perfil='Você já reconhece o passato prossimo e entende textos curtos com alguma ajuda.',
            objetivo='Ler por prazer e em quantidade: este é o coração da leitura extensiva.',
            como='Uma história por dia, ou a cada dois dias. Não pare a cada palavra: se entendeu a ideia, siga em frente. Ao terminar, escreva o riassunto em três frases.',
            meta='Pelo menos 15 das 22 histórias em 3 a 5 semanas.',
            avancar='Você ouve uma história nova acompanhando o texto e entende quase tudo sem pausar.'),
    3: dict(nome='Il ponte', pt='A ponte', nivel='A2–B1',
            curta='Textos mais longos e as primeiras adaptações de clássicos (Collodi, O. Henry, Maupassant).',
            perfil='Você lê histórias A1–A2 com facilidade e quer enredos mais ricos.',
            objetivo='Ganhar fôlego para textos de 15 a 30 minutos e se acostumar com o imperfetto na narração.',
            como='Divida em duas sessões se precisar. Sublinhe só as palavras que se repetem: são essas que vale a pena aprender.',
            meta='As 9 histórias em 3 a 4 semanas.',
            avancar='Você termina um texto de 20 minutos sem perder o fio da história.'),
    4: dict(nome='Il sentiero', pt='O caminho', nivel='B1–B2',
            curta='Adaptações de Verga, Maupassant, Poe, De Amicis, Tarchetti e Sacchetti: a literatura já aparece por trás do texto.',
            perfil='Você entende o essencial de textos sobre temas conhecidos e quer vocabulário mais rico.',
            objetivo='Preparar-se para a literatura original: estilo, ironia, descrições longas.',
            como='Leia a adaptação e, se gostar, abra o original no catálogo (Rosso Malpelo e La roba estão nos contos de Verga). Comparar as duas versões é um ótimo exercício.',
            meta='As 15 histórias em 4 a 6 semanas.',
            avancar='Os textos B1–B2 ficaram confortáveis e você sente vontade de ler algo "de verdade".'),
    5: dict(nome='Primi originali', pt='Primeiros originais', nivel='B1+ · estimativa',
            curta='Literatura infantojuvenil completa, sem adaptação: Pinocchio, Cuore, Gian Burrasca, Storie allegre.',
            perfil='Você lê B1 com conforto, mas ainda não leu um livro inteiro em italiano.',
            objetivo='Ler o primeiro livro inteiro em italiano original e aprender a reconhecer o passato remoto.',
            como='Escolha um livro e leia um capítulo por dia. Aceite não entender tudo: acompanhar a história já basta. Quando houver audiolivro, ouça junto nos capítulos difíceis.',
            meta='1 ou 2 livros em 2 a 3 meses.',
            avancar='Você terminou um livro original e reconhece as formas do passato remoto (fu, ebbe, disse, fece).'),
    6: dict(nome='La grande letteratura', pt='A grande literatura', nivel='B2+ · estimativa',
            curta='Contos e romances de 1850 a 1930: Verga, Pirandello, Deledda, Svevo, Serao, Capuana e muitos outros.',
            perfil='Você já leu um livro original e quer explorar autores e estilos.',
            objetivo='Leitura extensiva de literatura: variar autores e alternar contos curtos com romances longos.',
            como='Comece pelos contos curtos de Capuana, De Roberto e Albertazzi: 90 deles podem ser baixados aqui mesmo, em PDF. Depois alterne um romance com três ou quatro contos.',
            meta='Sem meta fixa: esta etapa dura o tempo que você quiser.',
            avancar='A sintaxe do século XIX não assusta mais e você quer encarar os grandes clássicos.'),
    7: dict(nome='I classici', pt='Os clássicos', nivel='C1+ · estimativa',
            curta='Dante, Petrarca, Boccaccio, Maquiavel, Leopardi, Foscolo e Pascoli: italiano antigo ou de estilo muito elaborado.',
            perfil='Você lê literatura do século XIX com conforto.',
            objetivo='Conhecer as obras que fundaram a língua e a literatura italiana.',
            como='Leia com uma edição comentada ou uma tradução em português ao lado. Trechos bastam: algumas novelle do Decameron, alguns cantos do Inferno.',
            meta='Um clássico por vez, aos poucos.',
            avancar='Não há próxima etapa. Buona lettura!'),
}
for t in TAPPE: TAPPE[t]['obras'] = [w for w in W if w['tappa'] == t]

def find(title, author=None):
    for w in W:
        if w['titulo'] == title and (author is None or w['autor'] == author): return w
    sys.exit(f'Obra não encontrada no catálogo: {title}')

def picks(t):
    if t == 6:
        per, out = {}, []
        for w in sorted((w for w in W if w['tappa'] == 6 and w['pdf']), key=lambda w: w['palavras']):
            if w['palavras'] > 1500 and per.get(w['autor'], 0) < 2:
                per[w['autor']] = per.get(w['autor'], 0) + 1; out.append(w)
        out.sort(key=lambda w: w['palavras'])
        romances = [find(*x) for x in [('Storia di una capinera', 'Giovanni Verga'), ('I Malavoglia', 'Giovanni Verga'),
                    ('Canne al vento', 'Grazia Deledda'), ('Il fu Mattia Pascal', 'Luigi Pirandello'),
                    ('Le tigri di Mompracem', 'Emilio Salgari'), ('La coscienza di Zeno', 'Italo Svevo'),
                    ('Tutte le novelle', 'Giovanni Verga'), ('I promessi sposi', 'Alessandro Manzoni')]]
        return [('Contos curtos para começar (PDF disponível aqui)', out), ('Romances e coletâneas essenciais', romances)]
    if t == 7:
        order = ['Operette morali', 'Il Principe', 'La mandragola - La Clizia - Belfagor', 'Ultime lettere di Jacopo Ortis', 'I Canti',
                 'Myricae', 'Decameron', 'Il Trecentonovelle', 'Canzoniere', 'La Divina Commedia di Dante']
        return [('Uma ordem possível', [find(x) for x in order])]
    return [('Leituras desta etapa', TAPPE[t]['obras'])]

def dur(m):
    if m < 60: return f'~{m} min'
    h = m / 60
    return f'~{round(h)} h' if h >= 2 else f'~{h:.1f} h'.replace('.', ',')

def autor_line(w):
    if w['adaptado_de']: return f"Adaptação de {w['adaptado_de']}"
    if w['autor']: return w['autor']
    return 'História do próprio site'

def main_link(w):
    return SITE + w['pdf'] if w['pdf'] else w['url']

ICON_AUDIO = '<svg aria-hidden="true" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round"><path d="M3 14v-2a9 9 0 0 1 18 0v2"/><rect x="3" y="14" width="4" height="7" rx="1.5"/><rect x="17" y="14" width="4" height="7" rx="1.5"/></svg>'
ICON_DL = '<svg aria-hidden="true" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round"><path d="M12 3v12m0 0-5-5m5 5 5-5M4 21h16"/></svg>'

# ───────────────────────── layout ─────────────────────────
NAV = [('./', 'Início'), ('catalogo.html', 'Catálogo'), ('plano.html', 'Plano de leitura'), ('guia.html', 'Como usar'), ('sobre.html', 'Sobre e direitos')]

def page(fname, title, desc, body, extra_head=''):
    nav = ''.join(f'<a href="{h}"{" aria-current=\"page\"" if h == fname or (h == './' and fname == 'index.html') else ""}>{t}</a>' for h, t in NAV)
    canon = SITE + ('' if fname == 'index.html' else fname)
    # A 404 é servida em qualquer caminho inexistente: base absoluta para os recursos, sem canonical e fora do índice.
    seo = ('<base href="/biblioteca-italiana/">\n<meta name="robots" content="noindex">' if fname == '404.html'
           else f'<link rel="canonical" href="{canon}">\n<meta property="og:url" content="{canon}">')
    full_title = title if fname == 'index.html' else f'{title} · Biblioteca Italiana'
    return f'''<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(full_title)}</title>
<meta name="description" content="{e(desc)}">
{seo}
<meta property="og:type" content="website">
<meta property="og:site_name" content="Biblioteca Italiana">
<meta property="og:title" content="{e(full_title)}">
<meta property="og:description" content="{e(desc)}">
<meta property="og:image" content="{SITE}assets/og.png">
<meta property="og:locale" content="pt_BR">
<meta name="twitter:card" content="summary_large_image">
<meta name="theme-color" content="#a63d2f">
<link rel="icon" href="assets/favicon.svg" type="image/svg+xml">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Fraunces:ital,opsz,wght@0,9..144,500;0,9..144,600;1,9..144,500;1,9..144,600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="assets/style.css?v={VER}">
{extra_head}</head>
<body>
<a class="skip" href="#conteudo">Pular para o conteúdo</a>
<header class="site-header"><div class="wrap">
<a class="brand" href="./"><img src="assets/favicon.svg" alt="" width="34" height="34"><span>Biblioteca Italiana<small>leggere per imparare</small></span></a>
<button class="nav-toggle" type="button" id="nav-toggle" aria-expanded="false" aria-controls="menu" hidden>Menu</button>
<nav class="nav" id="menu" aria-label="Principal">{nav}</nav>
</div></header>
<main id="conteudo">
{body}
</main>
<footer class="site-footer"><div class="wrap">
<div><p><b>Biblioteca Italiana</b> · leituras em italiano, do A1 aos clássicos.</p>
<p>Projeto pessoal e sem fins lucrativos, feito por um brasileiro que estuda italiano. Não é afiliado às fontes citadas.</p></div>
<div><p><a href="{REPO}">Código e dados no GitHub</a> · <a href="{REPO}/issues">Avisar um problema</a></p>
<p>Textos do projeto: CC BY 4.0 · Atualizado em {TODAY}</p></div>
</div></footer>
<script src="assets/app.js?v={VER}" defer></script>
</body>
</html>
'''

def card(w):
    t = w['tappa']; tags = [f'<span class="tag t">Tappa {t} · {e(w["nivel"] if w["nivel"] != "Original" else "Original")}</span>']
    if w['audio']: tags.append(f'<span class="tag audio">{ICON_AUDIO} {"Audiolivro" if w["audio_tipo"] == "audiolivro" else "Áudio"}</span>')
    if w['pdf']: tags.append(f'<span class="tag dl-tag">{ICON_DL} PDF aqui</span>')
    acts = []
    if w['pdf']:
        acts.append(f'<a class="btn small accent" href="{e(w["pdf"])}" download>Baixar PDF</a>')
        acts.append(f'<a class="btn small ghost" href="{e(w["url"])}" target="_blank" rel="noopener">Original (TEI) ↗</a>')
    elif w['url_alternativa']:
        acts.append(f'<a class="btn small" href="{e(w["url_alternativa"])}" target="_blank" rel="noopener">Procurar no Liber Liber ↗</a>')
        acts.append(f'<a class="btn small ghost" href="{e(w["url"])}" target="_blank" rel="noopener">Texto TEI ↗</a>')
    else:
        lab = 'Ler e ouvir na fonte ↗' if w['audio_tipo'] == 'na página da história' else 'Abrir na fonte ↗'
        acts.append(f'<a class="btn small" href="{e(w["url"])}" target="_blank" rel="noopener">{lab}</a>')
    if w['audio_tipo'] == 'audiolivro':
        acts.append(f'<a class="btn small ghost" href="{e(w["audio_url"])}" target="_blank" rel="noopener">Audiolivro ↗</a>')
    acts.append('<button type="button" class="mark want" data-kind="quero" aria-pressed="false">Quero ler</button>')
    acts.append('<button type="button" class="mark" data-kind="lido" aria-pressed="false">✓ Lido</button>')
    meta = ' · '.join(x for x in [w['forma'], dur(w['minutos']), f"{w['paginas']} págs.", w['fonte']] if x)
    q = ' '.join(str(x) for x in [w['titulo'], w['autor'], w['adaptado_de'], w['fonte'], w['forma'], w['nivel'], ' '.join(w['temas'])] if x)
    import unicodedata
    qn = unicodedata.normalize('NFD', q.lower()); qn = ''.join(c for c in qn if unicodedata.category(c) != 'Mn')
    verify = '<p class="verify">Direitos em verificação: só link.</p>' if w['verificar'] else ''
    return (f'<article class="obra c{t}" id="{e(w["id"])}" data-i="{w["i"]}" data-tappa="{t}" data-forma="{e(w["forma"])}" '
            f'data-fonte="{e(w["fonte"])}" data-audio="{int(w["audio"])}" data-pdf="{int(bool(w["pdf"]))}" data-min="{w["minutos"]}" '
            f'data-t="{e(w["titulo"])}" data-q="{e(qn)}">'
            f'<div class="top">{"".join(tags)}</div><h3>{e(w["titulo"])}</h3><p class="autor">{e(autor_line(w))}</p>'
            f'<p class="meta">{e(meta)}</p>{verify}<div class="acts">{"".join(acts)}</div></article>')

def tappa_cards(link=True):
    out = []
    for t, d in TAPPE.items():
        n = len(d['obras']); a = sum(w['audio'] for w in d['obras'])
        inner = (f'<span class="num">Tappa {t}</span><h3>{e(d["nome"])}</h3><span class="lvl">{e(d["pt"])} · {e(d["nivel"])}</span>'
                 f'<p>{e(d["curta"])}</p><div class="count">{n} obras{f" · {a} com áudio" if a else ""}</div>')
        out.append(f'<a class="tappa c{t}" href="catalogo.html?tappa={t}">{inner}</a>' if link else f'<div class="tappa c{t}">{inner}</div>')
    return '<div class="tappe">' + ''.join(out) + '</div>'

# ───────────────────────── páginas ─────────────────────────
N = len(W); NA = sum(w['audio'] for w in W); NP = sum(bool(w['pdf']) for w in W)

def p_index():
    ld = json.dumps({'@context': 'https://schema.org', '@type': 'WebSite', 'name': 'Biblioteca Italiana', 'url': SITE, 'inLanguage': 'pt-BR',
                     'description': 'Catálogo gratuito de leituras em italiano organizadas por nível, do A1 à literatura clássica.'}, ensure_ascii=False)
    body = f'''<section class="hero"><div class="wrap">
<p class="kicker">Benvenuti! Una biblioteca per chi impara l'italiano.</p>
<h1>Leia italiano no seu nível.</h1>
<p class="lead">{N} leituras gratuitas, do A1 aos clássicos, reunidas de seis acervos confiáveis e organizadas em sete etapas. Sem cadastro e sem precisar saber nada de GitHub.</p>
<ul class="stats"><li><b>{N}</b>obras</li><li><b>{NA}</b>com áudio</li><li><b>{NP}</b>PDFs para baixar aqui</li><li><b>7</b>etapas de leitura</li></ul>
<div class="cta"><a class="btn accent" href="catalogo.html">Explorar o catálogo</a><a class="btn ghost" href="plano.html#onde-comecar">Não sei meu nível</a></div>
</div></section>
<section class="wrap"><h2>O percurso, em sete <i>tappe</i></h2>
<p>Em italiano, <i>tappa</i> é uma etapa de viagem. Escolha a sua pelo nível que você tem hoje; cada cartão abre o catálogo já filtrado.</p>
{tappa_cards()}
</section>
<section class="wrap"><h2>Como funciona</h2>
<ol class="steps">
<li><b>Escolha uma leitura.</b><br>Filtre por etapa, por áudio ou por duração. Ou deixe o catálogo sortear uma.</li>
<li><b>Leia na fonte original.</b><br>Cada obra leva à página oficial do acervo. Os 90 contos marcados como "PDF aqui" baixam direto deste site.</li>
<li><b>Ouça, anote, resuma.</b><br>Leia ouvindo o áudio, anote até 5 palavras e escreva três frases sobre a história.</li>
<li><b>Marque seu progresso.</b><br>Use "Quero ler" e "Lido" no site ou a planilha de acompanhamento.</li>
</ol></section>
<section class="wrap"><h2>Para baixar</h2>
<div class="downloads">
<div class="dl"><h3>Plano de leitura (PDF)</h3><p>O percurso completo: método, as sete etapas, quanto ler, o que ler e quando avançar, com dicas para brasileiros.</p><a class="btn accent" href="downloads/{PDF_NAME}" download>{ICON_DL} Baixar o plano</a> <a class="btn ghost" href="plano.html">Ler online</a></div>
<div class="dl"><h3>Diário de leitura (planilha)</h3><p>Todas as {N} obras com links, status, datas e notas. Tem painel de progresso, registro de sessões, vocabulário com revisão e resumos.</p><a class="btn accent" href="downloads/{XLSX_NAME}" download>{ICON_DL} Baixar a planilha</a> <a class="btn ghost" href="guia.html#planilha">Como usar</a></div>
</div>
<p class="note">Os livros não ficam guardados aqui: quase sempre o link leva ao acervo que publica a obra. Só hospedamos arquivos com licença que permite redistribuir. <a href="sobre.html#direitos">Entenda por quê.</a></p>
</section>'''
    return page('index.html', 'Biblioteca Italiana — leituras em italiano por nível, do A1 aos clássicos',
                f'{N} leituras gratuitas em italiano organizadas por nível (A1 a C1): histórias graduadas com áudio, contos e clássicos. Com plano de leitura em PDF e planilha de acompanhamento.',
                body, f'<script type="application/ld+json">{ld}</script>\n')

def p_catalogo():
    fontes = sorted({w['fonte'] for w in W})
    formas = ['História graduada', 'Conto', 'Contos (coletânea)', 'Romance', 'Poesia', 'Teatro', 'Prosa e ensaio']
    chips_t = ''.join(f'<button type="button" class="chip-btn c{t}" data-group="tappa" data-value="{t}" aria-pressed="false">{t} · {e(d["nome"])}</button>' for t, d in TAPPE.items())
    chips_f = ''.join(f'<button type="button" class="chip-btn" data-group="forma" data-value="{e(f)}" aria-pressed="false">{e(f)}</button>' for f in formas)
    body = f'''<div class="wrap"><h1>Catálogo</h1>
<p>{N} obras em ordem sugerida de leitura, da mais fácil à mais difícil. Combine os filtros à vontade; o endereço da página guarda a busca para você compartilhar.</p></div>
<div class="filters"><div class="wrap">
<div class="search">
<label class="sr" for="q" hidden>Buscar</label>
<input id="q" type="search" placeholder="Buscar título, autor ou tema (ex.: Verga, mistero, Pinocchio)" autocomplete="off">
<select id="fonte" aria-label="Fonte"><option value="">Todas as fontes</option>{"".join(f'<option>{e(f)}</option>' for f in fontes)}</select>
<select id="ordem" aria-label="Ordenar"><option value="sugerida">Ordem sugerida</option><option value="curtas">Mais curtas primeiro</option><option value="longas">Mais longas primeiro</option><option value="titulo">Título (A–Z)</option></select>
</div>
<button type="button" class="filters-toggle" id="filtros-botao" aria-expanded="true" aria-controls="filtros-mais" hidden>Mais filtros</button>
<div class="filters-more" id="filtros-mais">
<div class="chips" role="group" aria-label="Etapa">{chips_t}</div>
<div class="chips" role="group" aria-label="Forma">{chips_f}</div>
<div class="toggles">
<label><input type="checkbox" id="f-audio"> Com áudio</label>
<label><input type="checkbox" id="f-pdf"> PDF para baixar aqui</label>
<label>Minha lista: <select id="lista"><option value="">todas</option><option value="quero">quero ler</option><option value="lido">já li</option><option value="nao-lido">ainda não li</option></select></label>
</div></div></div></div>
<div class="wrap">
<div class="result-bar"><span id="count" aria-live="polite">{N} obras</span>
<span><button type="button" class="btn small ghost" id="sorteio">🎲 Sortear uma leitura</button> <button type="button" class="btn small ghost" id="limpar">Limpar filtros</button></span></div>
<div class="grid" id="grid">{"".join(card(w) for w in W)}</div>
<p class="empty" id="empty" hidden>Nenhuma obra com esses filtros. Tente limpar algum deles.</p>
<p class="note">"Quero ler" e "Lido" ficam salvos só neste navegador, sem cadastro. O tempo é uma estimativa para quem está aprendendo. Os níveis das tappe 5 a 7 são estimativas nossas; os das tappe 1 a 4 são os declarados pelas fontes. Também é possível baixar o catálogo como <a href="dados/catalogo.csv">CSV</a> ou <a href="dados/catalogo.json">JSON</a>.</p>
</div>'''
    return page('catalogo.html', 'Catálogo de leituras em italiano por nível',
                f'Catálogo com {N} leituras em italiano: filtre por nível (A1 a C1), áudio, duração e forma. Histórias graduadas, contos e romances clássicos com links para as fontes oficiais.', body)

def plan_body(pdf=False):
    """Conteúdo do plano de leitura, usado no site e no PDF."""
    base = SITE if pdf else ''
    link = lambda w: f'<a href="{e(main_link(w) if pdf else (w["pdf"] or w["url"]))}">{e(w["titulo"])}</a>'
    def table(ws):
        rows = ''.join(f'<tr><td>{link(w)}</td><td>{e(autor_line(w))}</td><td>{e(w["nivel"])}</td><td>{dur(w["minutos"])}</td>'
                       f'<td>{"áudio" if w["audio_tipo"] == "na página da história" else ("audiolivro" if w["audio"] else "—")}</td>'
                       f'<td>{"PDF no site" if w["pdf"] else e(w["fonte"])}</td></tr>' for w in ws)
        return f'<div class="table-scroll"><table><thead><tr><th>Obra</th><th>Autor</th><th>Nível</th><th>Tempo</th><th>Áudio</th><th>Onde</th></tr></thead><tbody>{rows}</tbody></table></div>'
    tappe_html = ''
    for t, d in TAPPE.items():
        groups = ''.join(f'<h3>{e(h)}</h3>{table(ws)}' for h, ws in picks(t))
        more = '' if pdf else f'<p><a href="catalogo.html?tappa={t}">Ver as {len(d["obras"])} obras desta etapa no catálogo →</a></p>'
        if pdf and t == 6: more = f'<p>As {len(d["obras"])} obras desta etapa estão no catálogo: {SITE}catalogo.html?tappa=6</p>'
        tappe_html += f'''<section class="tappa-sec c{t}" id="tappa-{t}">
<div class="tappa-head"><h2>Tappa {t} · <i>{e(d["nome"])}</i></h2><span class="tag t">{e(d["pt"])} · {e(d["nivel"])}</span></div>
<p>{e(d["curta"])}</p>
<div class="two">
<div class="box"><b>Para quem</b><br>{e(d["perfil"])}<br><br><b>Objetivo</b><br>{e(d["objetivo"])}</div>
<div class="box"><b>Como ler</b><br>{e(d["como"])}<br><br><b>Meta</b><br>{e(d["meta"])}</div>
</div>
{groups}
<p class="note"><b>Sinal para avançar:</b> {e(d["avancar"])}</p>{more}
</section>'''
    return f'''
<p class="lead">Um caminho prático para sair do A1 e chegar à literatura italiana original, usando só as obras deste catálogo. Com 20 a 30 minutos por dia, as etapas 1 a 4 levam de 3 a 5 meses.</p>

<h2 id="onde-comecar">Onde começar?</h2>
<p>Faça este teste de cinco minutos. Leia o primeiro parágrafo de cada texto abaixo, em ordem, e pare no primeiro em que você precisar do dicionário mais de uma vez por frase.</p>
<ol>
<li>{link(find("La sveglia rotta"))} (A1). Se foi difícil, comece pela <b>Tappa 1</b>.</li>
<li>{link(find("Goffredo va al mercato"))} (A1–A2). Se foi difícil, comece pela <b>Tappa 1</b>; se foi tranquilo, pela <b>Tappa 2</b>.</li>
<li>{link(find("Il mistero di Fastigliano"))} (A2–B1). Se foi tranquilo, comece pela <b>Tappa 3</b>.</li>
<li>{link(find("Rosso Malpelo"))} (B1–B2). Se foi tranquilo, vá para a <b>Tappa 4</b>, ou direto à <b>Tappa 5</b> se já lê com facilidade.</li>
</ol>
<p>Na dúvida, comece uma etapa abaixo. Ler fácil demais custa pouco; ler difícil demais faz desistir.</p>

<h2 id="metodo">O método: quatro hábitos</h2>
<div class="two">
<div class="box"><b>1. Leia fácil e leia muito</b><br>Quase todo o seu tempo deve ir para textos que você entende quase inteiros. É a <i>leitura extensiva</i>: o vocabulário se fixa pela repetição, não pelo dicionário. <b>Regra prática:</b> se uma página tem mais de 5 a 8 palavras desconhecidas que atrapalham a história, o texto está difícil demais. Volte uma etapa sem culpa.</div>
<div class="box"><b>2. Ouça o que você lê</b><br>Para os textos com áudio, use três passagens: <b>(1)</b> leia sem áudio; <b>(2)</b> leia ouvindo; <b>(3)</b> só ouça, de olhos fechados ou no ônibus. Repetir a terceira passagem em outros dias treina o ouvido de graça.</div>
<div class="box"><b>3. Colete poucas palavras, mas bem</b><br>No máximo <b>cinco palavras por história</b>, sempre com a frase em que apareceram. Revise depois de 1, 3 e 7 dias; a planilha calcula as datas para você. Prefira palavras que se repetem no texto.</div>
<div class="box"><b>4. Escreva o <i>riassunto in tre frasi</i></b><br>Ao terminar, escreva três frases em italiano: <b>Chi?</b> (quem e onde) · <b>Cosa succede?</b> (o que acontece) · <b>Come finisce?</b> (como termina, ou o que você achou). Exemplo inventado: <i>Marco perde il treno per Roma. Aspetta due ore alla stazione e conosce una ragazza simpatica. Alla fine è contento di aver perso il treno.</i></div>
</div>

<h3>Uma semana típica (30 minutos por dia)</h3>
<div class="table-scroll keep"><table><thead><tr><th>Dia</th><th>O que fazer</th></tr></thead><tbody>
<tr><td>Segunda a sexta</td><td>20 min de leitura nova + 10 min ouvindo o áudio de um texto já lido, ou revisando palavras</td></tr>
<tr><td>Sábado</td><td>Escrever os <i>riassunti</i> da semana e revisar o vocabulário</td></tr>
<tr><td>Domingo</td><td>Livre: releia a história favorita da semana, só ouvindo</td></tr>
</tbody></table></div>

<h3>Ferramentas</h3>
<p>A <b>planilha de acompanhamento</b> ({base}downloads/{XLSX_NAME}) tem todas as obras deste plano, com status, datas, nota, registro de sessões, vocabulário com datas de revisão e espaço para os <i>riassunti</i>. O <b>catálogo online</b> ({SITE}catalogo.html) permite filtrar e sortear leituras.</p>

<h2 id="etapas">As sete etapas</h2>
<p>Os níveis das etapas 1 a 4 são os declarados pelas próprias fontes. Os das etapas 5 a 7 (obras originais, sem adaptação) são estimativas nossas.</p>
{tappe_html}

<h2 id="dicas">Dicas para brasileiros</h2>
<h3>Falsos amigos (<i>falsi amici</i>) que aparecem muito</h3>
<div class="table-scroll keep"><table><thead><tr><th>Italiano</th><th>Significa</th><th>Não é</th></tr></thead><tbody>
<tr><td>burro</td><td>manteiga</td><td>burro (asino)</td></tr>
<tr><td>caldo</td><td>quente</td><td>caldo (brodo)</td></tr>
<tr><td>camera</td><td>quarto</td><td>câmera (macchina fotografica)</td></tr>
<tr><td>guardare</td><td>olhar</td><td>guardar (mettere via, conservare)</td></tr>
<tr><td>salire</td><td>subir</td><td>sair (uscire)</td></tr>
<tr><td>fermare</td><td>parar</td><td>firmar (firmare)</td></tr>
<tr><td>squisito</td><td>delicioso</td><td>esquisito (strano)</td></tr>
<tr><td>morbido</td><td>macio</td><td>mórbido (morboso)</td></tr>
<tr><td>cena</td><td>jantar</td><td>cena (scena)</td></tr>
<tr><td>prestare</td><td>emprestar</td><td>prestar (servire, valere)</td></tr>
<tr><td>topo</td><td>rato</td><td>topo (cima)</td></tr>
<tr><td>subito</td><td>imediatamente</td><td>súbito (improvviso)</td></tr>
</tbody></table></div>
<h3>O <i>passato remoto</i>: o tempo verbal da literatura</h3>
<p>A partir da Tappa 5 quase toda narração está no <i>passato remoto</i>. Basta reconhecê-lo, sem precisar usá-lo. As formas mais frequentes:</p>
<div class="table-scroll keep"><table><thead><tr><th>Infinitivo</th><th>ele/ela</th><th>eles/elas</th><th>Português</th></tr></thead><tbody>
<tr><td>essere</td><td>fu</td><td>furono</td><td>foi / foram</td></tr>
<tr><td>avere</td><td>ebbe</td><td>ebbero</td><td>teve / tiveram</td></tr>
<tr><td>dire</td><td>disse</td><td>dissero</td><td>disse / disseram</td></tr>
<tr><td>fare</td><td>fece</td><td>fecero</td><td>fez / fizeram</td></tr>
<tr><td>venire</td><td>venne</td><td>vennero</td><td>veio / vieram</td></tr>
<tr><td>vedere</td><td>vide</td><td>videro</td><td>viu / viram</td></tr>
<tr><td>prendere</td><td>prese</td><td>presero</td><td>pegou / pegaram</td></tr>
<tr><td>parlare · credere · partire</td><td>parlò · credette · partì</td><td>parlarono · credettero · partirono</td><td>regulares em -ò, -é/-ette, -ì</td></tr>
</tbody></table></div>
<h3>Italiano do século XIX: o que muda</h3>
<p>Nos textos antigos (Tappe 6 e 7), algumas formas parecem estranhas mas são só grafias da época: <i>egli / ella</i> = lui / lei; <i>ne' · co' · pe'</i> = nei · coi · per i; <i>giuoco, figliuolo</i> = gioco, figliolo; <i>dì</i> = giorno. Os acentos também podem aparecer diferentes (<i>perchè</i> em vez de <i>perché</i>).</p>

'''

def p_plano():
    body = f'''<div class="wrap narrow prose">
<h1>Plano de leitura</h1>
<p><a class="btn accent" href="downloads/{PDF_NAME}" download>{ICON_DL} Baixar em PDF</a> <a class="btn ghost" href="downloads/{XLSX_NAME}" download>{ICON_DL} Planilha de acompanhamento</a></p>
{plan_body(False)}
</div>'''
    return page('plano.html', 'Plano de leitura em italiano, do A1 à literatura',
                'Plano de leitura gratuito para aprender italiano lendo: método de leitura extensiva com áudio, sete etapas do A1 ao C1, obras recomendadas e dicas para brasileiros. PDF para baixar.', body)

def p_guia():
    body = f'''<div class="wrap narrow prose">
<h1>Como usar</h1>
<p class="lead">Você não precisa de conta, cadastro nem de saber o que é GitHub. Tudo funciona no navegador, inclusive no celular.</p>

<h2>Em 3 minutos</h2>
<ol>
<li>Descubra seu ponto de partida com o <a href="plano.html#onde-comecar">teste rápido do plano</a>.</li>
<li>Abra o <a href="catalogo.html">catálogo</a> e clique na sua etapa (<i>tappa</i>).</li>
<li>Escolha uma obra e clique em <b>Abrir na fonte</b> (ou <b>Baixar PDF</b>, quando houver).</li>
<li>Ao terminar, clique em <b>✓ Lido</b>. Pronto, o site lembra para você.</li>
</ol>

<h2>O catálogo</h2>
<ul>
<li><b>Etapas (1 a 7):</b> do mais fácil ao mais difícil. Dá para marcar várias ao mesmo tempo.</li>
<li><b>Forma:</b> história graduada (escrita para estudantes), conto, coletânea, romance, poesia, teatro.</li>
<li><b>Com áudio:</b> histórias com gravação na página da fonte, ou clássicos com audiolivro.</li>
<li><b>PDF para baixar aqui:</b> os 90 contos que podemos hospedar legalmente.</li>
<li><b>Sortear uma leitura:</b> escolhe uma obra ao acaso entre as filtradas que você ainda não leu.</li>
<li><b>Minha lista:</b> "Quero ler" e "Lido" ficam salvos <i>neste navegador</i>. Se limpar os dados do navegador ou trocar de aparelho, a lista some. Para um registro permanente, use a planilha.</li>
<li><b>Compartilhar uma busca:</b> o endereço da página muda com os filtros. Copie e envie (ex.: <a href="catalogo.html?tappa=2&amp;audio=1">tappa 2 com áudio</a>).</li>
</ul>

<h2>Lendo no celular ou no tablet</h2>
<ul>
<li><b>PDFs</b> abrem direto no navegador. Para guardar, use "Compartilhar → Salvar em Arquivos" (iPhone) ou "Baixar" (Android).</li>
<li><b>Quick Italian e Learn Italian Online:</b> a história e o player de áudio estão na própria página. Leia ali mesmo, com fones.</li>
<li><b>Project Gutenberg:</b> prefira o formato <b>EPUB</b>, que abre no Apple Books, no Google Play Livros ou em leitores como o Kobo. Para o Kindle, use a opção "Kindle".</li>
<li><b>Liber Liber:</b> oferece PDF, EPUB e, em vários clássicos, audiolivro gratuito.</li>
<li><b>Wikisource:</b> dá para ler online ou baixar em EPUB/PDF pelo botão de download da página.</li>
</ul>

<h2 id="planilha">A planilha de acompanhamento</h2>
<p><a class="btn accent" href="downloads/{XLSX_NAME}" download>{ICON_DL} Baixar a planilha</a></p>
<p>Abre no Excel, no LibreOffice ou no <b>Google Planilhas</b> (em <i>Arquivo → Importar → Fazer upload</i>; ou envie para o Google Drive e abra com o Planilhas). Tem seis abas:</p>
<ul>
<li><b>Comece aqui:</b> instruções curtas.</li>
<li><b>Painel:</b> seu progresso por etapa, minutos lidos, palavras aprendidas. Calcula sozinho.</li>
<li><b>Catálogo:</b> as {N} obras com link. Escolha o <b>Status</b> (Quero ler, Lendo, Lido, Abandonei) e preencha datas e nota.</li>
<li><b>Sessões:</b> anote cada sessão de leitura (data, obra, minutos). Serve para ver sua constância.</li>
<li><b>Vocabulário:</b> palavra, tradução e a frase do texto. As datas de revisão (1, 3 e 7 dias) aparecem sozinhas.</li>
<li><b>Riassunti:</b> seus resumos em três frases.</li>
</ul>

<h2>Perguntas frequentes</h2>
<details><summary>É gratuito?</summary><p>Sim. Todas as obras são de acesso gratuito nas fontes, e o site, o plano e a planilha também são gratuitos.</p></details>
<details><summary>Por que a maioria das obras é só um link?</summary><p>Porque acesso gratuito não é o mesmo que permissão para redistribuir. Só hospedamos arquivos cuja licença permite isso; para o resto, levamos você à página oficial. Veja <a href="sobre.html#direitos">a política de direitos</a>.</p></details>
<details><summary>O que significam A1, A2, B1, B2, C1?</summary><p>São os níveis do Quadro Europeu Comum de Referência (QECR/CEFR). A1 e A2 são o básico, B1 e B2 o intermediário, C1 e C2 o avançado. Nas etapas 1 a 4 o nível é o declarado pelas fontes; nas etapas 5 a 7, como são obras originais, o nível é uma estimativa nossa.</p></details>
<details><summary>Quanto tempo leva cada leitura?</summary><p>O tempo do catálogo é uma estimativa para quem está aprendendo, calculada pelo número de palavras e pela etapa. Leia no seu ritmo.</p></details>
<details><summary>Encontrei um link quebrado ou um erro.</summary><p>Avise abrindo uma <a href="{REPO}/issues">issue no GitHub</a> (precisa de uma conta gratuita) ou pelo canal indicado no repositório.</p></details>
<details><summary>Posso sugerir uma obra ou uma fonte?</summary><p>Sim, pelo mesmo caminho. Fontes novas só entram com licença clara.</p></details>
</div>'''
    return page('guia.html', 'Como usar', 'Guia rápido: como usar o catálogo de leituras em italiano, ler no celular, usar a planilha de acompanhamento e perguntas frequentes.', body)

def p_sobre():
    counts = {f: sum(1 for w in W if w['fonte'] == f) for f in sorted({w['fonte'] for w in W})}
    src_rows = [
        ('Quick Italian', 'https://quickitalian.eu', 'Histórias graduadas A1–A2, com áudio', 'Direitos do site; sem licença aberta localizada', 'Link para a página da história'),
        ('Learn Italian Online', 'https://learn-italian-online.net', 'Histórias graduadas A1–C1, algumas com áudio', 'Maioria "© learn-italian-online.net"; algumas histórias A1 com aviso CC BY-NC-ND', 'Link para a página da história'),
        ('CLiGS textbox', 'https://github.com/cligs/textbox', 'Corpus acadêmico de contos e romances (1830–1930) em TEI', 'Contos: CC BY-SA 4.0 · Romances: CC BY-NC-SA 4.0', 'Contos: PDF hospedado aqui · Romances: link'),
        ('Liber Liber', 'https://liberliber.it', 'Clássicos em edição digital, com audiolivros', 'CC BY-NC-SA 4.0 (por obra)', 'Link para a página da obra'),
        ('Project Gutenberg', 'https://www.gutenberg.org', 'Obras em domínio público, em EPUB e HTML', 'Domínio público nos EUA + licença Project Gutenberg', 'Link para a página da obra'),
        ('Wikisource', 'https://it.wikisource.org', 'Transcrições revisadas de clássicos', 'CC BY-SA 4.0', 'Link para a página da obra'),
    ]
    rows = ''.join(f'<tr><td><a href="{u}">{n}</a><br><small>{counts.get(n.replace(" textbox", ""), 0)} obras</small></td><td>{c}</td><td>{l}</td><td>{m}</td></tr>' for n, u, c, l, m in src_rows)
    body = f'''<div class="wrap narrow prose">
<h1>Sobre e direitos</h1>
<p class="lead">A Biblioteca Italiana nasceu de uma frustração comum: existem ótimas leituras gratuitas em italiano, mas estão espalhadas por dezenas de sites. Este projeto reúne, organiza por nível e indica um caminho.</p>
<p>É um projeto pessoal e sem fins lucrativos de um brasileiro que estuda italiano. Não tem anúncios, não coleta dados e não exige cadastro. Não é afiliado a nenhuma das fontes.</p>

<h2>As fontes</h2>
<div class="table-scroll"><table><thead><tr><th>Fonte</th><th>O que oferece</th><th>Licença</th><th>Como aparece aqui</th></tr></thead><tbody>{rows}</tbody></table></div>
<p>Todo o mérito das obras é de seus autores e dos acervos que as publicam. Se você gosta de uma fonte, visite-a e, se puder, apoie-a diretamente.</p>

<h2 id="direitos">Política de direitos autorais</h2>
<ul>
<li><b>Acesso gratuito não é permissão.</b> Um PDF baixável de graça continua protegido se a fonte não der uma licença explícita. Por isso, a maior parte do catálogo é formada por <b>links</b> para as páginas oficiais.</li>
<li><b>Só hospedamos o que a licença permite sem condições comerciais:</b> os 90 contos do corpus CLiGS, licenciados em CC BY-SA 4.0 (a licença está declarada em cada arquivo TEI original). Os PDFs foram gerados a partir desse TEI, sem alterar o texto, e são distribuídos sob a mesma licença CC BY-SA 4.0, com crédito ao CLiGS.</li>
<li><b>Domínio público do original não cobre adaptações.</b> Uma versão simplificada de um conto de Verga pertence a quem a escreveu. Por isso, as histórias graduadas aparecem sempre como link.</li>
<li><b>Texto e áudio têm direitos separados.</b> Os áudios ficam nas páginas das fontes; não hospedamos nenhum.</li>
<li><b>Na dúvida, link.</b> Três obras com dúvidas sobre direitos nos EUA aparecem marcadas como "em verificação", só com link. Uma adaptação de obra ainda protegida foi retirada do catálogo público.</li>
</ul>
<p class="note warn">É titular de direitos e quer corrigir ou retirar algo? Abra uma <a href="{REPO}/issues">issue no GitHub</a>. A correção é feita com prioridade.</p>

<h2>Licença deste projeto</h2>
<ul>
<li><b>Textos do site, plano de leitura, planilha e catálogo (dados):</b> CC BY 4.0. Pode usar e adaptar, com crédito.</li>
<li><b>Código (HTML, CSS, JS e scripts):</b> MIT.</li>
<li><b>PDFs dos contos CLiGS:</b> CC BY-SA 4.0, com crédito a <a href="https://github.com/cligs/textbox">CLiGS textbox</a>.</li>
<li><b>As obras listadas com link:</b> seguem a licença de cada fonte.</li>
</ul>
</div>'''
    return page('sobre.html', 'Sobre o projeto e direitos autorais', 'Sobre a Biblioteca Italiana: fontes, licenças e a política de direitos autorais do catálogo de leituras em italiano.', body)

def p_404():
    return page('404.html', 'Página não encontrada', 'Página não encontrada.',
                '<div class="wrap narrow"><h1>Pagina non trovata</h1><p>Essa página não existe. Que tal voltar ao <a href="./">início</a> ou ao <a href="catalogo.html">catálogo</a>?</p></div>')

# ───────────────────────── PDF do plano ─────────────────────────
PDF_CSS = '''
@page { size: A4; margin: 18mm 17mm 20mm; @bottom-center { content: "Biblioteca Italiana · Plano de leitura · " counter(page); font: 8.5pt Helvetica, Arial, sans-serif; color: #7a7f80; } }
@page :first { margin: 0; @bottom-center { content: none; } }
body { font: 10.4pt/1.5 Georgia, "Iowan Old Style", serif; color: #1f2a2e; }
a { color: #a63d2f; text-decoration: none; }
h1, h2, h3 { font-family: Georgia, serif; line-height: 1.2; }
h2 { font-size: 17pt; margin: 18pt 0 6pt; border-top: 1px solid #ddd2bd; padding-top: 10pt; break-after: avoid; }
h3 { font-size: 12pt; margin: 12pt 0 4pt; break-after: avoid; }
p, li { orphans: 3; widows: 3; }
.lead { font-size: 12pt; color: #4a5559; }
table { width: 100%; border-collapse: collapse; font: 8.6pt/1.35 Helvetica, Arial, sans-serif; margin: 4pt 0 10pt; }
th { text-align: left; font-size: 7.4pt; text-transform: uppercase; letter-spacing: .04em; color: #6b7275; border-bottom: 1.2px solid #1f2a2e; padding: 4pt 5pt; }
td { border-bottom: .5px solid #ddd2bd; padding: 3.5pt 5pt; vertical-align: top; }
tr { break-inside: avoid; }
.two { display: flex; gap: 8pt; flex-wrap: wrap; }
.two > .box { flex: 1 1 45%; }
.box { border: .8px solid #ddd2bd; border-radius: 6pt; padding: 8pt 10pt; margin: 5pt 0; background: #fbf8f1; font-size: 9.6pt; break-inside: avoid; }
.note { background: #f3ede1; border-left: 3pt solid #5b7a5e; padding: 6pt 10pt; font-size: 9.6pt; break-inside: avoid; }
.tag { font: bold 8pt Helvetica, Arial, sans-serif; color: #fff; background: var(--c); padding: 2pt 7pt; border-radius: 9pt; }
.tappa-head { display: flex; align-items: baseline; gap: 8pt; border-top: 3pt solid var(--c); margin-top: 6pt; }
.tappa-head h2 { border: 0; color: var(--c); margin-top: 8pt; }
.tappa-sec { break-before: page; }
td:nth-child(3), td:nth-child(4), td:nth-child(5) { white-space: nowrap; }
.keep table, table.keep { break-inside: avoid; } .keep { break-inside: avoid; }
.c1 { --c: #4f8a5b; } .c2 { --c: #2f8a86; } .c3 { --c: #2f6f9f; } .c4 { --c: #5a58a6; } .c5 { --c: #b8862a; } .c6 { --c: #b5532f; } .c7 { --c: #8a2f4a; }
.cover { width: 210mm; height: 297mm; background: #f6f1e7; padding: 34mm 22mm; box-sizing: border-box; position: relative; page-break-after: always; }
.cover .k { font: italic 14pt Georgia, serif; color: #a63d2f; }
.cover h1 { font-size: 40pt; margin: 6pt 0 4pt; }
.cover .sub { font-size: 15pt; color: #4a5559; max-width: 140mm; }
.cover .bars { display: flex; gap: 3mm; margin: 22mm 0 10mm; }
.cover .bars div { flex: 1; height: 9mm; border-radius: 2mm; }
.cover ol { font: 11pt/1.8 Helvetica, Arial, sans-serif; padding-left: 5mm; }
.cover .foot { position: absolute; bottom: 22mm; left: 22mm; right: 22mm; font: 9pt/1.5 Helvetica, Arial, sans-serif; color: #6b7275; }
.cover img { width: 18mm; }
.table-scroll { overflow: visible; }
'''

def build_pdf():
    from weasyprint import HTML
    logo = open(os.path.join(DOCS, 'assets', 'favicon.svg'), encoding='utf-8').read()
    import base64
    logo_uri = 'data:image/svg+xml;base64,' + base64.b64encode(logo.encode()).decode()
    bars = ''.join(f'<div style="background:{c}"></div>' for c in ['#4f8a5b', '#2f8a86', '#2f6f9f', '#5a58a6', '#b8862a', '#b5532f', '#8a2f4a'])
    toc = ''.join(f'<li>Tappa {t} · <i>{e(d["nome"])}</i> — {e(d["pt"])} ({e(d["nivel"])})</li>' for t, d in TAPPE.items())
    doc = f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Plano de leitura · Biblioteca Italiana</title>
<style>{PDF_CSS}</style></head><body>
<section class="cover"><img src="{logo_uri}" alt=""><p class="k">Biblioteca Italiana · leggere per imparare</p>
<h1>Plano de leitura</h1><p class="sub">Do A1 à literatura italiana original, em sete etapas, com obras gratuitas reais, método e dicas para brasileiros.</p>
<div class="bars">{bars}</div><ol>{toc}</ol>
<p class="foot">Versão de {TODAY} · {SITE}<br>Texto sob licença CC BY 4.0. As obras citadas seguem as licenças de suas fontes; os links levam às páginas oficiais.</p></section>
{plan_body(True)}
</body></html>'''
    HTML(string=doc, base_url=DOCS).write_pdf(os.path.join(DOCS, 'downloads', PDF_NAME))

# ───────────────────────── planilha ─────────────────────────
def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    INK, ACC, PAPER, LINE = '1F2A2E', 'A63D2F', 'F6F1E7', 'DDD2BD'
    TC = {1: '4F8A5B', 2: '2F8A86', 3: '2F6F9F', 4: '5A58A6', 5: 'B8862A', 6: 'B5532F', 7: '8A2F4A'}
    head_font = Font(bold=True, color='FFFFFF'); head_fill = PatternFill('solid', fgColor=INK)
    thin = Border(bottom=Side(style='thin', color=LINE))
    wb = Workbook()

    def header(ws, cols, widths):
        for i, (c, wdt) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=1, column=i, value=c); cell.font = head_font; cell.fill = head_fill
            cell.alignment = Alignment(vertical='center', wrap_text=True); ws.column_dimensions[get_column_letter(i)].width = wdt
        ws.row_dimensions[1].height = 30; ws.freeze_panes = 'A2'

    # Comece aqui
    ws = wb.active; ws.title = 'Comece aqui'; ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 100
    lines = [('Biblioteca Italiana · Diário de leitura', Font(size=22, bold=True, color=ACC)),
             ('leggere per imparare', Font(size=12, italic=True, color='4A5559')), ('', None),
             ('Como usar, em quatro passos', Font(size=14, bold=True)),
             ('1. Na aba "Catálogo", encontre uma obra (use os filtros do cabeçalho: Tappa, Nível, Áudio...).', None),
             ('2. Clique em "abrir ↗" para ir à página oficial da obra. Ao começar, mude o Status para "Lendo"; ao terminar, para "Lido".', None),
             ('3. Registre cada sessão na aba "Sessões" e as palavras novas na aba "Vocabulário" (no máximo 5 por história).', None),
             ('4. Escreva seu resumo em três frases na aba "Riassunti". O "Painel" mostra seu progresso automaticamente.', None), ('', None),
             ('As etapas (tappe)', Font(size=14, bold=True))]
    lines += [(f'Tappa {t} · {d["nome"]} ({d["pt"]}, {d["nivel"]}): {d["curta"]}', None) for t, d in TAPPE.items()]
    lines += [('', None), ('Dicas', Font(size=14, bold=True)),
              ('• Se uma página tem mais de 5 a 8 palavras desconhecidas que atrapalham a história, desça uma etapa.', None),
              ('• Textos com áudio: leia; depois leia ouvindo; depois só ouça.', None),
              ('• O "Tempo (min)" é uma estimativa para quem está aprendendo.', None), ('', None),
              (f'Site, catálogo e plano de leitura: {SITE}', Font(color=ACC, underline='single')),
              ('Planilha sob licença CC BY 4.0. As obras seguem as licenças de suas fontes.', Font(size=9, color='6B7275'))]
    for r, (txt, f) in enumerate(lines, 2):
        c = ws.cell(row=r, column=2, value=txt); c.alignment = Alignment(wrap_text=True, vertical='top')
        if f: c.font = f
        if txt.startswith('Site,'): c.hyperlink = SITE
        if txt.startswith('Tappa'):
            t = int(txt.split()[1]); c.font = Font(color=TC[t])
    # Catálogo
    cat = wb.create_sheet('Catálogo')
    cols = ['Nº', 'Tappa', 'Nível', 'Título', 'Autor', 'Forma', 'Fonte', 'Tempo (min)', 'Áudio', 'Link', 'Status', 'Início', 'Fim', 'Nota (1–5)', 'Riassunto?', 'Comentários']
    header(cat, cols, [5, 7, 10, 38, 26, 17, 18, 10, 11, 9, 12, 11, 11, 9, 10, 40])
    for r, w in enumerate(W, 2):
        vals = [r - 1, w['tappa'], w['nivel'], w['titulo'], autor_line(w), w['forma'], w['fonte'], w['minutos'],
                {'audiolivro': 'Audiolivro', 'na página da história': 'Sim'}.get(w['audio_tipo'], ''), 'abrir ↗']
        for cidx, v in enumerate(vals, 1):
            c = cat.cell(row=r, column=cidx, value=v); c.border = thin
        link = cat.cell(row=r, column=10); link.hyperlink = main_link(w); link.font = Font(color=ACC, underline='single')
        t = cat.cell(row=r, column=2); t.fill = PatternFill('solid', fgColor=TC[w['tappa']]); t.font = Font(bold=True, color='FFFFFF'); t.alignment = Alignment(horizontal='center')
        for cidx in (12, 13): cat.cell(row=r, column=cidx).number_format = 'DD/MM/YYYY'
        for cidx in range(11, 17): cat.cell(row=r, column=cidx).border = thin
    last = len(W) + 1
    dv = DataValidation(type='list', formula1='"Quero ler,Lendo,Lido,Abandonei"', allow_blank=True); cat.add_data_validation(dv); dv.add(f'K2:K{last}')
    dvn = DataValidation(type='list', formula1='"1,2,3,4,5"', allow_blank=True); cat.add_data_validation(dvn); dvn.add(f'N2:N{last}')
    dvs = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); cat.add_data_validation(dvs); dvs.add(f'O2:O{last}')
    dvd = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                         error='Digite uma data (ex.: 15/03/2027).', errorTitle='Data'); cat.add_data_validation(dvd); dvd.add(f'L2:M{last}')
    for status, color in [('Lido', 'DCEBDD'), ('Lendo', 'FBEFD2'), ('Quero ler', 'E3ECF5'), ('Abandonei', 'EEEEEE')]:
        cat.conditional_formatting.add(f'A2:P{last}', FormulaRule(formula=[f'$K2="{status}"'], fill=PatternFill('solid', fgColor=color)))
    cat.auto_filter.ref = f'A1:P{last}'
    # Sessões
    ses = wb.create_sheet('Sessões')
    header(ses, ['Data', 'Obra', 'Minutos', 'Páginas ou capítulos', 'Com áudio?', 'Observações'], [12, 40, 10, 18, 11, 50])
    obra_list = f"'Catálogo'!$D$2:$D${last}"
    for sh, col, rng in [(ses, 'B', 'B2:B1000')]:
        d = DataValidation(type='list', formula1=f'={obra_list}', allow_blank=True, showErrorMessage=False); sh.add_data_validation(d); d.add(rng)
    d = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); ses.add_data_validation(d); d.add('E2:E1000')
    for r in range(2, 1001): ses.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
    # Vocabulário
    voc = wb.create_sheet('Vocabulário')
    header(voc, ['Data', 'Parola', 'Tradução', 'Frase do texto', 'Obra', 'Revisar em (1 dia)', 'Revisar em (3 dias)', 'Revisar em (7 dias)', 'Aprendida?'],
           [12, 18, 20, 50, 32, 13, 13, 13, 11])
    d = DataValidation(type='list', formula1=f'={obra_list}', allow_blank=True, showErrorMessage=False); voc.add_data_validation(d); d.add('E2:E2000')
    d = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); voc.add_data_validation(d); d.add('I2:I2000')
    for r in range(2, 501):
        voc.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
        for cidx, n in ((6, 1), (7, 3), (8, 7)):
            c = voc.cell(row=r, column=cidx, value=f'=IF(A{r}="","",A{r}+{n})'); c.number_format = 'DD/MM/YYYY'; c.font = Font(color='6B7275')
    voc.conditional_formatting.add('F2:H500', FormulaRule(formula=['AND(F2<>"",F2<=TODAY(),$I2<>"Sim")'], fill=PatternFill('solid', fgColor='FBEFD2')))
    # Riassunti
    ria = wb.create_sheet('Riassunti')
    header(ria, ['Data', 'Obra', '1. Chi? (quem e onde)', '2. Cosa succede? (o que acontece)', '3. Come finisce? (fim ou opinião)'], [12, 32, 40, 48, 48])
    d = DataValidation(type='list', formula1=f'={obra_list}', allow_blank=True, showErrorMessage=False); ria.add_data_validation(d); d.add('B2:B1000')
    for r in range(2, 301):
        ria.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
        for cidx in (3, 4, 5): ria.cell(row=r, column=cidx).alignment = Alignment(wrap_text=True, vertical='top')
    # Painel (segunda aba)
    pn = wb.create_sheet('Painel', 1); pn.sheet_view.showGridLines = False
    for col, wdt in zip('ABCDEFGH', [9, 26, 9, 9, 9, 11, 12, 4]): pn.column_dimensions[col].width = wdt
    pn['A1'] = 'Seu progresso'; pn['A1'].font = Font(size=20, bold=True, color=ACC)
    pn['A2'] = 'Atualiza sozinho conforme você preenche as outras abas.'; pn['A2'].font = Font(italic=True, color='6B7275')
    kpis = [('Obras lidas', '=COUNTIF(Catálogo!$K:$K,"Lido")'), ('Lendo agora', '=COUNTIF(Catálogo!$K:$K,"Lendo")'),
            ('Lidas neste mês', '=COUNTIFS(Catálogo!$K:$K,"Lido",Catálogo!$M:$M,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))'),
            ('Minutos de leitura (sessões)', '=SUM(Sessões!$C:$C)'), ('Dias com leitura (sessões)', '=SUMPRODUCT((Sessões!$A$2:$A$1000<>"")/COUNTIF(Sessões!$A$2:$A$1000,Sessões!$A$2:$A$1000&""))'),
            ('Palavras anotadas', '=COUNTA(Vocabulário!$B$2:$B$2000)'), ('Palavras aprendidas', '=COUNTIF(Vocabulário!$I:$I,"Sim")'),
            ('Riassunti escritos', '=COUNTA(Riassunti!$B$2:$B$1000)')]
    for i, (lab, f) in enumerate(kpis):
        r = 4 + i
        pn.cell(row=r, column=1, value=lab).font = Font(color='4A5559'); pn.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = pn.cell(row=r, column=3, value=f); c.font = Font(size=13, bold=True); c.alignment = Alignment(horizontal='right')
    r0 = 4 + len(kpis) + 1
    for i, h in enumerate(['Tappa', 'Etapa', 'Obras', 'Lidas', 'Lendo', 'Quero ler', '% lidas'], 1):
        c = pn.cell(row=r0, column=i, value=h); c.font = head_font; c.fill = head_fill
    for k, (t, d) in enumerate(TAPPE.items(), 1):
        r = r0 + k
        pn.cell(row=r, column=1, value=t).fill = PatternFill('solid', fgColor=TC[t]); pn.cell(row=r, column=1).font = Font(bold=True, color='FFFFFF')
        pn.cell(row=r, column=1).alignment = Alignment(horizontal='center')
        pn.cell(row=r, column=2, value=f'{d["nome"]} ({d["nivel"]})')
        pn.cell(row=r, column=3, value=f'=COUNTIF(Catálogo!$B:$B,A{r})')
        pn.cell(row=r, column=4, value=f'=COUNTIFS(Catálogo!$B:$B,A{r},Catálogo!$K:$K,"Lido")')
        pn.cell(row=r, column=5, value=f'=COUNTIFS(Catálogo!$B:$B,A{r},Catálogo!$K:$K,"Lendo")')
        pn.cell(row=r, column=6, value=f'=COUNTIFS(Catálogo!$B:$B,A{r},Catálogo!$K:$K,"Quero ler")')
        c = pn.cell(row=r, column=7, value=f'=IF(C{r}=0,0,D{r}/C{r})'); c.number_format = '0%'
        for cidx in range(1, 8): pn.cell(row=r, column=cidx).border = thin
    rt = r0 + 8
    pn.cell(row=rt, column=2, value='Total').font = Font(bold=True)
    for cidx, col in ((3, 'C'), (4, 'D'), (5, 'E'), (6, 'F')):
        pn.cell(row=rt, column=cidx, value=f'=SUM({col}{r0+1}:{col}{r0+7})').font = Font(bold=True)
    c = pn.cell(row=rt, column=7, value=f'=IF(C{rt}=0,0,D{rt}/C{rt})'); c.number_format = '0%'; c.font = Font(bold=True)
    pn.conditional_formatting.add(f'G{r0+1}:G{rt}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='5B7A5E'))
    wb.save(os.path.join(DOCS, 'downloads', XLSX_NAME))

# ───────────────────────── dados abertos, sitemap ─────────────────────────
def build_data():
    os.makedirs(os.path.join(DOCS, 'dados'), exist_ok=True)
    pub = [{k: v for k, v in w.items() if k != 'i'} for w in W]
    json.dump(pub, open(os.path.join(DOCS, 'dados', 'catalogo.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    cols = ['tappa', 'nivel', 'titulo', 'autor', 'adaptado_de', 'forma', 'temas', 'fonte', 'minutos', 'paginas', 'palavras', 'audio_tipo',
            'url', 'audio_url', 'pdf', 'licenca']
    with open(os.path.join(DOCS, 'dados', 'catalogo.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        wr = csv.writer(f); wr.writerow(cols)
        for w in W:
            wr.writerow([', '.join(w[c]) if c == 'temas' else (SITE + w[c] if c == 'pdf' and w[c] else w[c]) for c in cols])
    urls = ['', 'catalogo.html', 'plano.html', 'guia.html', 'sobre.html']
    sm = ''.join(f'<url><loc>{SITE}{u}</loc><lastmod>{TODAY}</lastmod></url>' for u in urls)
    open(os.path.join(DOCS, 'sitemap.xml'), 'w').write(f'<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">{sm}</urlset>\n')
    open(os.path.join(DOCS, 'robots.txt'), 'w').write(f'User-agent: *\nAllow: /\nSitemap: {SITE}sitemap.xml\n')
    open(os.path.join(DOCS, '.nojekyll'), 'w').write('')

if __name__ == '__main__':
    os.makedirs(os.path.join(DOCS, 'downloads'), exist_ok=True)
    for fname, fn in [('index.html', p_index), ('catalogo.html', p_catalogo), ('plano.html', p_plano), ('guia.html', p_guia),
                      ('sobre.html', p_sobre), ('404.html', p_404)]:
        open(os.path.join(DOCS, fname), 'w', encoding='utf-8').write(fn())
    build_data()
    build_xlsx()
    if '--sem-pdf' not in sys.argv: build_pdf()
    print(f'ok: {N} obras · {NA} com áudio · {NP} PDFs hospedados')
